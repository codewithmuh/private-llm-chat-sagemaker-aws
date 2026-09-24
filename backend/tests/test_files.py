"""Uploads, text extraction, serving files, OCR."""

from __future__ import annotations

import io

import docx
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from PIL import Image

from apps.files.models import Attachment

from .conftest import Client

pytestmark = pytest.mark.django_db


def text_pdf(text: str) -> bytes:
    """A minimal one-page PDF with a real text layer, built by hand."""
    stream = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def png(color="white", size=(40, 20)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


def upload(client: Client, name: str, data: bytes):
    return client.post("/api/files/", {"file": SimpleUploadedFile(name, data)}, format="multipart")


def test_text_pdf(auth_client: Client):
    res = upload(auth_client, "report.pdf", text_pdf("Quarterly revenue grew"))
    assert res.status_code == 201, res.content
    body = res.json()
    assert body["kind"] == "document" and body["page_count"] == 1
    assert body["has_text"] is True
    assert "Quarterly revenue grew" in auth_client.get(f"/api/files/{body['id']}/text/").json()["text"]


def test_scanned_pdf_has_no_text(auth_client: Client):
    buf = io.BytesIO()
    Image.new("RGB", (200, 200), "white").save(buf, format="PDF")
    body = upload(auth_client, "scan.pdf", buf.getvalue()).json()
    assert body["page_count"] == 1
    assert body["has_text"] is False


def test_docx(auth_client: Client):
    document = docx.Document()
    document.add_paragraph("Meeting notes")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Owner", "Ada"
    buf = io.BytesIO()
    document.save(buf)
    body = upload(auth_client, "notes.docx", buf.getvalue()).json()
    text = auth_client.get(f"/api/files/{body['id']}/text/").json()["text"]
    assert "Meeting notes" in text and "Owner | Ada" in text


def test_image_upload_and_content(auth_client: Client):
    body = upload(auth_client, "photo.png", png()).json()
    assert body["kind"] == "image" and body["content_type"] == "image/png"
    res = auth_client.get(body["url"])
    assert res.status_code == 200
    assert res["Content-Type"] == "image/png"
    assert b"".join(res.streaming_content).startswith(b"\x89PNG")


def test_html_is_served_as_download(auth_client: Client):
    body = upload(auth_client, "page.html", b"<script>alert(1)</script>").json()
    res = auth_client.get(body["url"])
    assert res["Content-Type"] == "application/octet-stream"
    assert res["Content-Disposition"].startswith("attachment")


def test_rejects_unsupported_and_fake_files(auth_client: Client):
    assert upload(auth_client, "tool.exe", b"MZ").status_code == 415
    assert upload(auth_client, "fake.png", b"not an image").status_code == 415


def test_size_limit(auth_client: Client, settings):
    settings.MAX_UPLOAD_MB = 1
    res = upload(auth_client, "big.txt", b"a" * (1024 * 1024 + 1))
    assert res.status_code == 413
    assert res.json()["code"] == "file_too_large"


def test_files_are_private(auth_client: Client, client_factory=None):
    body = upload(auth_client, "a.txt", b"hello").json()
    auth_client.post_json("/api/auth/logout/")
    assert auth_client.get(f"/api/files/{body['id']}/").status_code == 401


def test_ocr_with_a_vision_model(auth_client: Client, mock_model):
    body = upload(auth_client, "receipt.png", png()).json()
    res = auth_client.post_json(f"/api/files/{body['id']}/ocr/", {})
    assert res.status_code == 200, res.content
    data = res.json()
    assert data["engine"] == "llm:mock-chat"
    assert "mock OCR" in data["text"]
    assert data["attachment"]["ocr_done"] is True and data["attachment"]["has_text"] is True
    # Cached: the second call does not run again.
    assert auth_client.post_json(f"/api/files/{body['id']}/ocr/", {}).json()["duration_ms"] == 0


def test_ocr_without_any_engine(auth_client: Client, settings, monkeypatch):
    monkeypatch.setattr("apps.files.ocr.tesseract_available", lambda: False)
    body = upload(auth_client, "receipt.png", png()).json()
    res = auth_client.post_json(f"/api/files/{body['id']}/ocr/", {})
    assert res.status_code == 503


def test_delete_removes_the_stored_file(auth_client: Client):
    body = upload(auth_client, "a.txt", b"hello").json()
    attachment = Attachment.objects.get(pk=body["id"])
    storage, name = attachment.file.storage, attachment.file.name
    assert storage.exists(name)
    assert auth_client.delete(f"/api/files/{body['id']}/").status_code == 204
    assert not storage.exists(name)
