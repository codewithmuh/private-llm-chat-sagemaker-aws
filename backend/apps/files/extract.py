"""Getting text out of uploaded files.

A language model only reads text (and, if it has vision, images). So on
upload we pull the text out of every document once and store it; each chat
turn then pastes it into the prompt.

    .pdf               pypdfium2 (the PDF engine Chrome uses). Scanned PDFs
                       have no text layer: has_text stays false and the UI
                       offers OCR.
    .docx              python-docx: paragraphs and tables
    text / code        decoded as UTF-8 (falling back to Latin-1)
    images             validated with Pillow; text comes from OCR or vision
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from pathlib import PurePath

from PIL import Image, UnidentifiedImageError

IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}
DOCUMENT_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}
TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".json",
    ".jsonl",
    ".xml",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".log",
    ".html",
    ".htm",
    ".css",
    ".sql",
    ".sh",
    ".py",
    ".js",
    ".mjs",
    ".ts",
    ".tsx",
    ".jsx",
    ".java",
    ".kt",
    ".go",
    ".rs",
    ".c",
    ".h",
    ".cpp",
    ".hpp",
    ".cs",
    ".rb",
    ".php",
    ".swift",
    ".r",
    ".scala",
    ".tf",
    ".ipynb",
}
ACCEPTED_EXTENSIONS = set(IMAGE_TYPES) | set(DOCUMENT_TYPES) | TEXT_EXTENSIONS

# Keeps one giant upload from filling the database. ~150k tokens of text.
MAX_TEXT_CHARS = 600_000
# A PDF averaging fewer characters than this per page is treated as scanned.
MIN_CHARS_PER_PAGE = 20
# Refuse images that decompress into absurd sizes ("decompression bombs").
Image.MAX_IMAGE_PIXELS = 60_000_000


class UnsupportedFile(Exception):
    pass


@dataclass
class Extracted:
    kind: str  # "image" | "document"
    content_type: str
    text: str = ""
    page_count: int | None = None
    width: int | None = None
    height: int | None = None
    error: str = ""


def extension(filename: str) -> str:
    return PurePath(filename or "").suffix.lower()


def inspect_file(filename: str, data: bytes) -> Extracted:
    """Classify an upload and extract what we can. Raises UnsupportedFile."""
    ext = extension(filename)
    if ext not in ACCEPTED_EXTENSIONS:
        raise UnsupportedFile(f"{ext or 'This file type'} is not supported.")

    if ext in IMAGE_TYPES:
        try:
            with Image.open(io.BytesIO(data)) as img:
                img.verify()  # cheap integrity check without decoding every pixel
            with Image.open(io.BytesIO(data)) as img:
                width, height = img.size
        except (UnidentifiedImageError, OSError, Image.DecompressionBombError) as exc:
            raise UnsupportedFile("This image could not be read.") from exc
        return Extracted(kind="image", content_type=IMAGE_TYPES[ext], width=width, height=height)

    if ext == ".pdf":
        return _pdf(data)
    if ext == ".docx":
        return _docx(data)
    return _plain_text(ext, data)


def _truncate(text: str) -> str:
    text = text.replace("\x00", "")
    if len(text) > MAX_TEXT_CHARS:
        return text[:MAX_TEXT_CHARS] + "\n\n[... truncated: the file is longer than the app keeps ...]"
    return text


def _pdf(data: bytes) -> Extracted:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as exc:
        raise UnsupportedFile("This PDF could not be opened (is it password protected?).") from exc
    pages: list[str] = []
    try:
        for index in range(len(pdf)):
            page = pdf[index]
            textpage = page.get_textpage()
            pages.append(textpage.get_text_bounded().strip())
            textpage.close()
            page.close()
        count = len(pdf)
    finally:
        pdf.close()
    joined = "\n\n".join(f"[Page {i + 1}]\n{t}" for i, t in enumerate(pages) if t)
    total = sum(len(t) for t in pages)
    scanned = count > 0 and total / count < MIN_CHARS_PER_PAGE
    return Extracted(
        kind="document",
        content_type=DOCUMENT_TYPES[".pdf"],
        text="" if scanned else _truncate(joined),
        page_count=count,
    )


def _docx(data: bytes) -> Extracted:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # python-docx raises several unrelated types
        raise UnsupportedFile("This Word document could not be opened.") from exc
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return Extracted(kind="document", content_type=DOCUMENT_TYPES[".docx"], text=_truncate("\n".join(parts)))


def _plain_text(ext: str, data: bytes) -> Extracted:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    if ext == ".ipynb":
        text = _notebook_text(text)
    return Extracted(kind="document", content_type="text/plain", text=_truncate(text))


def _notebook_text(raw: str) -> str:
    """Jupyter notebooks are JSON; keep just the cells' source."""
    try:
        nb = json.loads(raw)
    except ValueError:
        return raw
    cells = []
    for cell in nb.get("cells", []):
        source = "".join(cell.get("source", []))
        cells.append(f"# [{cell.get('cell_type', 'cell')}]\n{source}")
    return "\n\n".join(cells)
